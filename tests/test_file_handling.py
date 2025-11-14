from app.main import get_extracted_text
import pytest
import pymupdf
from typing import Optional
from io import BytesIO
from docx import Document
from fastapi import HTTPException


def make_pdf(text: Optional[str] = None) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    if text:
        page.insert_text((50, 50), text)
    b = doc.tobytes()
    doc.close()
    return b


def make_docx(text: Optional[str] = None) -> bytes:
    buf = BytesIO()
    doc = Document()
    if text:
        doc.add_paragraph(text)
    else:
        doc.add_paragraph("")
    doc.save(buf)
    return buf.getvalue()


@pytest.fixture
def docx_type():
    return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def test_pdf_extraction_success():
    pdf_bytes = make_pdf("Job info from PDF")
    result = get_extracted_text(pdf_bytes, "application/pdf", "resume.pdf")
    assert "Job info from PDF" in result


def test_pdf_extraction_invalid_binary():
    with pytest.raises(HTTPException):
        get_extracted_text(b"not a pdf", "application/pdf", "resume.pdf")


def test_pdf_extraction_empty_content():
    empty = make_pdf()
    with pytest.raises(HTTPException):
        get_extracted_text(empty, "application/pdf", "resume.pdf")


def test_docx_extraction_success(docx_type):
    docx_bytes = make_docx("Job info from DOCX")
    result = get_extracted_text(docx_bytes, docx_type, "resume.docx")
    assert "Job info from DOCX" in result


def test_docx_extraction_invalid(docx_type):
    with pytest.raises(HTTPException):
        get_extracted_text(b"not_docx", docx_type, "resume.docx")


def test_docx_extraction_empty(docx_type):
    empty = make_docx()
    with pytest.raises(HTTPException):
        get_extracted_text(empty, docx_type, "resume.docx")


def test_unsupported_format():
    with pytest.raises(HTTPException):
        get_extracted_text(b"123", "text/plain", "file.txt")
